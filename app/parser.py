from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook

from .config_io import load_resident_profile
from .constants import DEFAULT_CASE_CLASS, DEFAULT_SITE
from .utils import (
    case_year_from_date,
    extract_procedure_text,
    format_acgme_date,
    parse_principal_result_interpreter,
    parse_age_years,
    parse_role_metadata,
    parse_study_datetime,
    patient_type,
    patient_type_from_age,
    review_flags,
    row_hash,
)

REQUIRED_COLUMNS = [
    "Accession Number",
    "Exam Code",
    "Institution Name",
    "Modalities DICOM",
    "Modality",
    "Patient Birth Date",
    "Patient ID",
    "Principal Result Interpreter",
    "Report ID",
    "Report Snippet",
    "Study Date",
    "Study Description",
    "Study Instance UID",
]

MPOWER_REQUIRED_COLUMNS = [
    "Accession Number",
    "Modality",
    "Exam Code",
    "Exam Description",
    "CPT Code",
    "Report Text",
    "Patient Age",
    "Exam Started Date",
    "Report Finalized By",
]

MPOWER_HASH_COLUMNS = [c for c in MPOWER_REQUIRED_COLUMNS]

SECTION_STOP_RE = re.compile(
    r"^\s*(?:"
    r"Plan|PLAN|PROCEDURE DETAILS|PROCEDURE COMMENTS AND FINDINGS|Attestation|Additional Details|"
    r"Pre-procedure|Post-procedure diagnosis|Preoperative diagnosis|PREOPERATIVE DIAGNOSIS|"
    r"POST-OPERATIVE DIAGNOSIS|Indication|Additional clinical history|TECHNIQUE|FINDINGS|"
    r"COMPLICATIONS|CLINICAL HISTORY|SEDATION|ANESTHESIA|ACCESS/CLOSURE|Radiation Dose|Contrast|"
    r"[A-Z0-9 /()\-]+ PROCEDURE SUMMARY"
    r")\s*:?\s*$",
    re.I | re.M,
)

PERSONNEL_END_RE = re.compile(
    r"^\s*(?:"
    r"Pre-procedure diagnosis|Post-procedure diagnosis|Indication|Additional clinical history|"
    r"Adverse events|Complications|TECHNIQUE|PROCEDURE DETAILS|PROCEDURE COMMENTS|"
    r"Pre-procedure|FINDINGS|IMPRESSION|CLINICAL HISTORY|SEDATION|ANESTHESIA|ACCESS/CLOSURE"
    r")\b",
    re.I | re.M,
)

TITLE_LIKE_RE = re.compile(
    r"\b(?:GUIDED|ULTRASOUND|PARACENTESIS|THORACENTESIS|NEEDLE|BIOPSY|MYELO|ARTHROGRAM|"
    r"ANGIO|EMBOL|DRAIN|TUBE|CATHETER|PORT|FILTER|TIPS|THROMB|SCLEROTHERAPY)\b",
    re.I,
)

NOT_ATTEMPTED_RE = re.compile(
    r"\b(?:not attempted|not performed|no procedure was performed|procedure was not performed|deferred|unable to perform)\b",
    re.I,
)


def iter_raw_rows(xlsx_path: str | Path, sheet_name: str = "raw") -> Iterator[dict[str, Any]]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, "") for v in row):
            continue
        yield dict(zip(header, row))


def iter_mpower_raw_rows(csv_path: str | Path) -> Iterator[dict[str, Any]]:
    with Path(csv_path).open(newline="") as f:
        reader = csv.DictReader(f)
        header = reader.fieldnames or []
        missing = [c for c in MPOWER_REQUIRED_COLUMNS if c not in header]
        if missing:
            raise ValueError(f"Missing required mPower columns: {', '.join(missing)}")
        for row_number, row in enumerate(reader, start=2):
            if not any(v not in (None, "") for v in row.values()):
                continue
            row["__source_row_number"] = row_number
            yield row


def _nonblank_lines(text: str | None) -> list[str]:
    return [line.strip() for line in str(text or "").splitlines() if line.strip()]


def _clean_line(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def _strip_bullet(value: str) -> str:
    return _clean_line(re.sub(r"^\s*(?:[-*]|\d+[\.)])\s*", "", value))


def _is_heading(line: str) -> bool:
    clean = line.strip()
    if not clean:
        return False
    if SECTION_STOP_RE.match(clean):
        return True
    if re.match(r"^[A-Z][A-Z0-9 /()\-]{2,50}:?\s*$", clean):
        return True
    return False


def _section_after_heading(text: str, heading: str) -> str:
    pattern = re.compile(rf"^\s*{heading}\s*:?\s*$", re.I | re.M)
    match = pattern.search(text)
    if not match:
        inline = re.search(rf"^\s*{heading}\s*:\s*(.*?)$", text, re.I | re.M)
        if not inline:
            return ""
        start = inline.end()
        first = inline.group(1).strip()
    else:
        start = match.end()
        first = ""
    stop = SECTION_STOP_RE.search(text, start)
    body = text[start : stop.start() if stop else len(text)]
    lines = [first] if first else []
    lines.extend(_nonblank_lines(body))
    return "\n".join(line for line in lines if line)


def _extract_impression(text: str) -> str:
    return _section_after_heading(text, "IMPRESSION")


def _extract_fallback_findings(text: str) -> str:
    for heading in ("FINDINGS/IMPRESSION", "FINDINGS", "PROCEDURE COMMENTS AND FINDINGS"):
        value = _section_after_heading(text, heading)
        if value:
            return value
    return ""


def _extract_procedure_list(lines: list[str], start_index: int) -> list[str]:
    values: list[str] = []
    for line in lines[start_index:]:
        if re.match(r"^(?:Date(?: of service| of procedure| of procedure/surgery)?|EXAM DATE)\b", line, re.I):
            break
        if _is_heading(line) and not re.match(r"^\d+[\.)]\s+", line):
            break
        stripped = _strip_bullet(line)
        if stripped:
            values.append(stripped)
    return values


def _extract_procedure_title(text: str) -> tuple[str, list[str], str]:
    lines = _nonblank_lines(text)
    if not lines:
        return "", [], "missing"

    for idx, line in enumerate(lines):
        proc_match = re.match(r"^PROCEDURE\s*:\s*(.*)$", line, re.I)
        if proc_match:
            inline = proc_match.group(1).strip()
            if inline:
                return _clean_line(inline), [], "procedure_label"
            items = _extract_procedure_list(lines, idx + 1)
            return "; ".join(items), items, "procedure_label_block"

        procs_match = re.match(r"^PROCEDURES\s*:\s*(.*)$", line, re.I)
        if procs_match:
            inline = procs_match.group(1).strip()
            items = [p.strip() for p in re.split(r"\s*,\s*", inline) if p.strip()] if inline else []
            items.extend(_extract_procedure_list(lines, idx + 1))
            return "; ".join(items) if items else inline, items, "procedures_label"

    first = lines[0]
    if not re.match(r"^DATE\b", first, re.I) and TITLE_LIKE_RE.search(first):
        return _clean_line(first), [], "first_line_title"

    return _clean_line(first), [], "first_line_fallback"


def _extract_procedure_summaries(text: str) -> list[dict[str, Any]]:
    summaries: list[dict[str, Any]] = []
    pattern = re.compile(r"^\s*(.*PROCEDURE SUMMARY)\s*:\s*$", re.I | re.M)
    matches = list(pattern.finditer(text))
    for index, match in enumerate(matches):
        start = match.end()
        next_start = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        details = re.search(r"^\s*PROCEDURE DETAILS\s*:?\s*$", text[start:next_start], re.I | re.M)
        end = start + details.start() if details else next_start
        lines = _nonblank_lines(text[start:end])
        bullets = [_strip_bullet(line) for line in lines if re.match(r"^\s*[-*]\s+", line)]
        additional: list[str] = []
        for line in lines:
            m = re.match(r"^\s*-?\s*Additional procedure\(s\)\s*:\s*(.+)$", line, re.I)
            if m:
                value = _clean_line(m.group(1))
                if value and value.lower() not in {"none", "n/a", "na"}:
                    additional.extend(v.strip() for v in re.split(r"\s*,\s*", value) if v.strip())
        summaries.append(
            {
                "heading": _clean_line(match.group(1)),
                "lines": lines,
                "bullets": bullets,
                "additional_procedures": additional,
            }
        )
    return summaries


def _extract_personnel_section(text: str) -> str:
    heading = re.search(r"^\s*(?:Procedural Personnel|PROCEDURE PERSONNEL\s*:?)\s*$", text, re.I | re.M)
    if not heading:
        return ""
    start = heading.end()
    stop = PERSONNEL_END_RE.search(text, start)
    return text[start : stop.start() if stop else min(len(text), start + 1000)]


def _extract_label_value(section: str, label: str) -> str:
    m = re.search(rf"^\s*{label}\s*:\s*(.*)$", section, re.I | re.M)
    return _clean_line(m.group(1)) if m else ""


def _resident_block(text: str) -> str:
    m = re.search(r"^\s*RESIDENT\s*:\s*(.*)$", text, re.I | re.M)
    if not m:
        return ""
    first = _clean_line(m.group(1))
    lines = [first] if first else []
    for line in str(text[m.end() :]).splitlines():
        clean = line.strip()
        if not clean:
            continue
        if _is_heading(clean):
            break
        lines.append(clean)
    return "; ".join(lines)


def _split_people(value: str) -> list[str]:
    if not value or value.strip().lower() in {"none", "n/a", "na"}:
        return []
    normalized = re.sub(r"\b(?:MD|M\.D\.|DO|D\.O\.|PhD|MPH|MS|FSIR)\b\.?", "", value)
    normalized = re.sub(r"\bDr\.\s*", "", normalized, flags=re.I)
    parts = re.split(r"\s*;\s*|\s+\band\b\s+|\n+", normalized)
    if len(parts) == 1:
        parts = re.split(r"\s*,\s*(?=[A-Z][a-z]+(?:\s+[A-Z][a-z]+)*)", normalized)
    return [_clean_line(p).strip(",") for p in parts if _clean_line(p).strip(",")]


def _alias_match_positions(text: str, aliases: list[str]) -> list[int]:
    positions: list[int] = []
    for alias in aliases:
        pieces = alias.replace(",", " ").split()
        if not pieces:
            continue
        pattern = r"\b" + r"\s+".join(re.escape(piece) for piece in pieces) + r"\b"
        for match in re.finditer(pattern, text, re.I):
            positions.append(match.start())
    return sorted(set(positions))


def _resident_position(value: str, aliases: list[str]) -> int | None:
    if not value or value.strip().lower() in {"none", "n/a", "na"}:
        return None
    people = _split_people(value)
    if people:
        for idx, person in enumerate(people, start=1):
            if _alias_match_positions(person, aliases):
                return idx
    positions = _alias_match_positions(value, aliases)
    if not positions:
        return None
    first = positions[0]
    separators = [m.start() for m in re.finditer(r";|\band\b|,(?=\s*[A-Z])", value, re.I)]
    return 1 + sum(1 for s in separators if s < first)


def parse_mpower_role_metadata(report_text: str | None, aliases: list[str]) -> dict[str, Any]:
    fallback = {
        "role": "Primary",
        "role_confidence": "low",
        "resident_found_in_report": False,
        "resident_position": None,
        "role_parse_source": "fallback",
    }
    text = str(report_text or "")
    personnel = _extract_personnel_section(text)
    if personnel:
        resident_line = _extract_label_value(personnel, r"Resident physician\(s\)")
        position = _resident_position(resident_line, aliases)
        if position:
            return {
                "role": "Primary" if position == 1 else "Secondary",
                "role_confidence": "high",
                "resident_found_in_report": True,
                "resident_position": position,
                "role_parse_source": "personnel_resident_line",
            }

        other_line = _extract_label_value(personnel, "Other")
        position = _resident_position(other_line, aliases)
        if position:
            return {
                "role": "Primary" if position == 1 else "Secondary",
                "role_confidence": "high",
                "resident_found_in_report": True,
                "resident_position": position,
                "role_parse_source": "personnel_other_line",
            }

        non_resident = " ".join(
            value
            for value in [
                _extract_label_value(personnel, r"Fellow physician\(s\)"),
                _extract_label_value(personnel, r"Advanced practice provider\(s\)"),
            ]
            if value
        )
        position = _resident_position(non_resident, aliases)
        if position:
            return {
                "role": "Primary" if position == 1 else "Secondary",
                "role_confidence": "low",
                "resident_found_in_report": True,
                "resident_position": position,
                "role_parse_source": "non_resident_personnel_line",
            }

        return {**fallback, "role_parse_source": "personnel_no_resident_match"}

    resident_block = _resident_block(text)
    position = _resident_position(resident_block, aliases)
    if position:
        return {
            "role": "Primary" if position == 1 else "Secondary",
            "role_confidence": "high",
            "resident_found_in_report": True,
            "resident_position": position,
            "role_parse_source": "resident_heading",
        }

    if _alias_match_positions(text, aliases):
        return {**fallback, "resident_found_in_report": True, "resident_position": 1, "role_parse_source": "non_personnel_report_mention"}
    return fallback


def parse_mpower_report(report_text: str | None, aliases: list[str] | None = None) -> dict[str, Any]:
    text = str(report_text or "")
    aliases = aliases or []
    title, procedure_list, title_source = _extract_procedure_title(text)
    impression = _extract_impression(text)
    fallback_findings = "" if impression else _extract_fallback_findings(text)
    summaries = _extract_procedure_summaries(text)
    personnel_section = _extract_personnel_section(text)
    resident_block = _resident_block(text)
    personnel = {
        "attending": _extract_label_value(personnel_section, r"Attending physician\(s\)") or _extract_label_value(personnel_section, "Attending"),
        "fellow": _extract_label_value(personnel_section, r"Fellow physician\(s\)") or _extract_label_value(personnel_section, "FELLOW"),
        "resident": _extract_label_value(personnel_section, r"Resident physician\(s\)") or resident_block,
        "advanced_practice_provider": _extract_label_value(personnel_section, r"Advanced practice provider\(s\)") or _extract_label_value(personnel_section, "APP"),
        "other": _extract_label_value(personnel_section, "Other"),
        "section_found": bool(personnel_section),
    }

    candidate_phrases: list[str] = []
    for value in [title, *procedure_list, impression, fallback_findings]:
        candidate_phrases.extend(_nonblank_lines(value))
    for summary in summaries:
        candidate_phrases.extend(summary["bullets"])
        candidate_phrases.extend(summary["additional_procedures"])
    unique_candidates = list(dict.fromkeys(_strip_bullet(v) for v in candidate_phrases if _strip_bullet(v)))

    warnings: list[str] = []
    if not impression:
        warnings.append("missing_impression")
    if not summaries:
        warnings.append("missing_procedure_summary")
    if not personnel_section and not resident_block:
        warnings.append("missing_personnel_section")
    role_meta = parse_mpower_role_metadata(text, aliases)
    if role_meta["role_confidence"] == "low":
        warnings.append("personnel_role_low_confidence")
    if re.search(r"\baborted\b", text, re.I):
        warnings.append("aborted_language")
    if re.search(r"\bunsuccessful\b|not successful|failed attempt", text, re.I):
        warnings.append("unsuccessful_language")
    if NOT_ATTEMPTED_RE.search(text):
        warnings.append("not_attempted_language")

    return {
        "procedure_title": title,
        "procedure_title_source": title_source,
        "procedure_list": procedure_list,
        "impression": impression,
        "fallback_findings": fallback_findings,
        "procedure_summary_sections": summaries,
        "personnel": personnel,
        "candidate_procedure_phrases": unique_candidates,
        "parse_warnings": list(dict.fromkeys(warnings)),
    }


def parse_report_finalized_by(raw: str | None) -> str:
    if not raw:
        return ""
    text = str(raw).strip()
    if "," in text:
        last, first = [part.strip() for part in text.split(",", 1)]
        return f"{first.title()} {last.title()}".strip()
    return text


def transform_source_row(raw: dict[str, Any], resident_profile: dict[str, Any] | None = None) -> dict[str, Any]:
    profile = resident_profile or load_resident_profile()
    resident = profile["resident"]
    defaults = profile["defaults"]
    role_meta = parse_role_metadata(raw.get("Report Snippet"), resident.get("aliases", []))
    flags = review_flags(raw.get("Report Snippet"))
    study_dt = parse_study_datetime(raw["Study Date"])
    procedure_text = extract_procedure_text(raw.get("Report Snippet")) or str(raw.get("Study Description") or "")
    return {
        "accession_number": str(raw.get("Accession Number") or "").strip(),
        "study_date": study_dt.isoformat(),
        "patient_birth_date": str(raw.get("Patient Birth Date") or ""),
        "patient_age_years": None,
        "exam_code": str(raw.get("Exam Code") or "").strip(),
        "study_description": str(raw.get("Study Description") or "").strip(),
        "report_snippet": str(raw.get("Report Snippet") or ""),
        "procedure_text": procedure_text,
        "institution_name": str(raw.get("Institution Name") or "").strip(),
        "source_format": "visage_xlsx",
        "source_row_number": None,
        "modality": str(raw.get("Modality") or raw.get("Modalities DICOM") or "").strip(),
        "cpt_code": "",
        "duplicate_accession_flag": 0,
        "parsed_report_json": None,
        "principal_result_interpreter_raw": str(raw.get("Principal Result Interpreter") or "").strip(),
        "attending_name": parse_principal_result_interpreter(raw.get("Principal Result Interpreter")),
        "source_row_hash": row_hash(raw),
        "resident_found_in_report": int(bool(role_meta["resident_found_in_report"])),
        "resident_position": role_meta["resident_position"],
        "role_parse_source": role_meta["role_parse_source"],
        "aborted_flag": int(flags["aborted_flag"]),
        "unsuccessful_flag": int(flags["unsuccessful_flag"]),
        "no_procedure_flag": int(flags["no_procedure_flag"]),
        "needs_review_reason": flags["needs_review_reason"],
        "derived": {
            "case_id": str(raw.get("Accession Number") or "").strip(),
            "case_date": format_acgme_date(raw["Study Date"]),
            "case_year": case_year_from_date(
                raw["Study Date"],
                int(resident["expected_graduation_year"]),
                int(resident.get("pgy_max", 5)),
            ),
            "role": role_meta["role"],
            "role_confidence": role_meta["role_confidence"],
            "site": defaults.get("site", DEFAULT_SITE),
            "patient_type": patient_type(
                raw["Study Date"],
                raw.get("Patient Birth Date"),
                int(defaults.get("pediatric_age_cutoff", 18)),
            ),
            "case_class": defaults.get("case_class", DEFAULT_CASE_CLASS),
        },
    }


def transform_mpower_row(
    raw: dict[str, Any],
    duplicate_accession: bool = False,
    resident_profile: dict[str, Any] | None = None,
) -> dict[str, Any]:
    profile = resident_profile or load_resident_profile()
    resident = profile["resident"]
    defaults = profile["defaults"]
    study_dt = parse_study_datetime(raw["Exam Started Date"])
    parsed = parse_mpower_report(raw.get("Report Text"), resident.get("aliases", []))
    role_meta = parse_mpower_role_metadata(raw.get("Report Text"), resident.get("aliases", []))
    flags = review_flags(raw.get("Report Text"))
    warning_reasons = [w for w in parsed.get("parse_warnings", [])]
    if duplicate_accession:
        warning_reasons.append("duplicate_accession")
    base_reason = flags["needs_review_reason"]
    needs_review_reason = ", ".join(dict.fromkeys([r for r in [base_reason, *warning_reasons] if r]))
    age = parse_age_years(raw.get("Patient Age"))
    procedure_text = parsed.get("procedure_title") or str(raw.get("Exam Description") or "")
    return {
        "accession_number": str(raw.get("Accession Number") or "").strip(),
        "study_date": study_dt.isoformat(),
        "patient_birth_date": "",
        "patient_age_years": str(age) if age is not None else None,
        "exam_code": str(raw.get("Exam Code") or "").strip(),
        "study_description": str(raw.get("Exam Description") or "").strip(),
        "report_snippet": str(raw.get("Report Text") or ""),
        "procedure_text": procedure_text,
        "institution_name": "",
        "source_format": "mpower_csv",
        "source_row_number": raw.get("__source_row_number"),
        "modality": str(raw.get("Modality") or "").strip(),
        "cpt_code": str(raw.get("CPT Code") or "").strip(),
        "duplicate_accession_flag": int(duplicate_accession),
        "parsed_report_json": json.dumps(parsed, ensure_ascii=False, sort_keys=True),
        "principal_result_interpreter_raw": str(raw.get("Report Finalized By") or "").strip(),
        "attending_name": parse_report_finalized_by(raw.get("Report Finalized By")),
        "source_row_hash": row_hash(raw, MPOWER_HASH_COLUMNS),
        "resident_found_in_report": int(bool(role_meta["resident_found_in_report"])),
        "resident_position": role_meta["resident_position"],
        "role_parse_source": role_meta["role_parse_source"],
        "aborted_flag": int(flags["aborted_flag"] or "aborted_language" in warning_reasons),
        "unsuccessful_flag": int(flags["unsuccessful_flag"] or "unsuccessful_language" in warning_reasons),
        "no_procedure_flag": int(flags["no_procedure_flag"] or "not_attempted_language" in warning_reasons),
        "needs_review_reason": needs_review_reason,
        "derived": {
            "case_id": str(raw.get("Accession Number") or "").strip(),
            "case_date": format_acgme_date(raw["Exam Started Date"]),
            "case_year": case_year_from_date(
                raw["Exam Started Date"],
                int(resident["expected_graduation_year"]),
                int(resident.get("pgy_max", 5)),
            ),
            "role": role_meta["role"],
            "role_confidence": role_meta["role_confidence"],
            "site": defaults.get("site", DEFAULT_SITE),
            "patient_type": patient_type_from_age(
                raw.get("Patient Age"),
                int(defaults.get("pediatric_age_cutoff", 18)),
            ),
            "case_class": defaults.get("case_class", DEFAULT_CASE_CLASS),
        },
    }
